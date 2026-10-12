import copy
from datetime import date
from decimal import Decimal
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook
from screen import KINDS, calculate, export, read_roster

ASOF = date(2026, 10, 11)
ROSTER = [{'code': '000001', 'name': '=unsafe', 'date': '2026-09-30'}]


def ledger():
    events = []
    for kind, amount in zip(KINDS, [200, 30, 100]):
        e = dict(code='000001', kind=kind, event_id=kind, amount_cny=str(amount), date='2025-01-01', status='cancelled' if kind == 'buyback_cancelled' else 'completed', source='https://example.org/report', evidence='合成测试凭证')
        if kind == 'buyback_cancelled':
            e['cancellation_source'] = 'https://example.org/cancellation'
        if kind == 'equity_financing':
            e['method'] = 'IPO'
        events.append(e)
    return {'events': events, 'coverage': [dict(code='000001', kind=k, complete=True, first_listing_date='1991-01-01', through='2026-10-11', source='https://example.org/archive', scope='合成测试，同一发行人全部上市市场') for k in KINDS]}


def test_calculation_and_pending_exclusion():
    data = ledger()
    pending = copy.deepcopy(data['events'][1])
    pending.update(event_id='pending', status='pending', amount_cny='999')
    data['events'].append(pending)
    all_rows, chosen = calculate(ROSTER, data, ASOF)
    assert chosen.iloc[0]['净股东回报（元）'] == Decimal('130')
    assert chosen.iloc[0]['股东回报/融资比率'] == Decimal('2.3')


def test_missing_not_zero():
    data = ledger()
    data['coverage'].pop()
    rows, chosen = calculate(ROSTER, data, ASOF)
    assert chosen.empty
    assert rows.iloc[0]['累计股权融资（元）'] is None
    assert rows.iloc[0]['净股东回报（元）'] is None


def test_duplicate_rejected():
    data = ledger()
    data['events'].append(copy.deepcopy(data['events'][0]))
    with pytest.raises(ValueError, match='重复'):
        calculate(ROSTER, data, ASOF)


def test_cancellation_evidence_required():
    data = ledger()
    del data['events'][1]['cancellation_source']
    with pytest.raises(ValueError, match='注销凭证'):
        calculate(ROSTER, data, ASOF)


@pytest.mark.parametrize('method', ['debt', 'secondary_market'])
def test_invalid_financing(method):
    data = ledger()
    data['events'][2]['method'] = method
    with pytest.raises(ValueError, match='融资方式'):
        calculate(ROSTER, data, ASOF)


def test_zero_financing_and_cutoff():
    data = ledger()
    data['events'][2]['amount_cny'] = '0'
    data['events'][0]['date'] = '2027-01-01'
    rows, chosen = calculate(ROSTER, data, ASOF)
    assert chosen.iloc[0]['净股东回报（元）'] == Decimal('30')
    assert rows.iloc[0]['股东回报/融资比率'] is None


def test_roster_dates_and_codes():
    stream = BytesIO()
    pd.DataFrame({'日期': ['2026-09-30'] * 300, '成分券代码': [str(i) for i in range(1, 301)], '成分券名称': ['测试'] * 300}).to_excel(stream, index=False)
    stream.seek(0)
    result = read_roster(stream, ASOF)
    assert result[0]['code'] == '000001'
    assert result[0]['date'] == '2026-09-30'
    stream.seek(0)
    with pytest.raises(ValueError, match='晚于'):
        read_roster(stream, date(2026, 1, 1))


def test_export(tmp_path):
    data = ledger()
    rows, chosen = calculate(ROSTER, data, ASOF)
    path = tmp_path / 'result.xlsx'
    export(path, rows, chosen, [], ['合成数据'], data)
    book = load_workbook(path)
    assert len(book.sheetnames) == 6
    assert book['筛选结果']['A2'].value == '000001'
    assert book['筛选结果']['B2'].data_type == 's'
    assert book['筛选结果']['H2'].value == '130'


def test_negative_excluded_and_sorting():
    data = ledger()
    other = {'code': '000002', 'name': 'test', 'date': '2026-09-30'}
    for e in list(data['events']):
        item = dict(e, code='000002')
        if item['kind'] == 'dividend':
            item['amount_cny'] = '300'
        data['events'].append(item)
    data['coverage'] += [dict(c, code='000002') for c in list(data['coverage'])]
    _, chosen = calculate(ROSTER + [other], data, ASOF)
    assert chosen['股票代码'].tolist() == ['000002', '000001']
    data['events'][0]['amount_cny'] = '1'
    _, chosen = calculate(ROSTER, data, ASOF)
    assert chosen.empty


def test_inconsistent_scope_rejected():
    data = ledger()
    data['coverage'][0]['first_listing_date'] = '2000-01-01'
    with pytest.raises(ValueError, match='一致'):
        calculate(ROSTER, data, ASOF)


def test_unknown_status_rejected():
    data = ledger()
    data['events'][0]['status'] = 'done_typo'
    with pytest.raises(ValueError, match='状态'):
        calculate(ROSTER, data, ASOF)


def test_official_bilingual_roster():
    stream = BytesIO()
    pd.DataFrame({'日期Date': [20261009] * 300, '成份券代码Constituent Code': list(range(1, 301)), '成份券名称Constituent Name': ['测试'] * 300}).to_excel(stream, index=False)
    stream.seek(0)
    result = read_roster(stream, ASOF)
    assert result[0]['date'] == '2026-10-09'
    assert result[0]['code'] == '000001'


def test_initial_ipo_before_listing_included():
    data = ledger()
    data['events'][2]['date'] = '1990-12-01'
    _, chosen = calculate(ROSTER, data, ASOF)
    assert chosen.iloc[0]['累计股权融资（元）'] == Decimal('100')
