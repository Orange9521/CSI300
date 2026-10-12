"""Supplement the workbook with original CNINFO PDFs, never infer lifetime coverage."""
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from pypdf import PdfReader
import requests
from fetch_public import read_json, dump, write_sheet
from screen import read_roster

QUERY = 'https://www.cninfo.com.cn/new/hisAnnouncement/query'
MASTER = 'https://www.cninfo.com.cn/new/data/szse_stock.json'


def query_batch(stocks, keyword, period, folder):
    params = {'pageNum': 1, 'pageSize': 100, 'column': 'szse', 'tabName': 'fulltext',
        'stock': ';'.join(f'{s["code"]},{s["orgId"]}' for s in stocks), 'searchkey': keyword,
        'secid': '', 'plate': '', 'category': '', 'trade': '', 'seDate': period,
        'sortName': 'time', 'sortType': 'desc', 'isHLtitle': 'true'}
    found, expected, seen = [], None, set()
    page = 1
    while True:
        params['pageNum'] = page
        # This endpoint honors pageNum in the query string; body-only paging
        # can silently repeat page 1. Check uniqueness as well as total count.
        r = requests.post(QUERY, params={'pageNum': page}, data=params, timeout=35)
        r.raise_for_status()
        payload = r.json()
        if not isinstance(payload.get('totalAnnouncement'), int):
            raise ValueError('公告查询未返回合法总数')
        if expected is None:
            expected = payload['totalAnnouncement']
        elif payload['totalAnnouncement'] != expected:
            raise ValueError('公告分页过程中数量改变')
        (folder / f'query_{stocks[0]["code"]}_{keyword}_{page}.json').write_bytes(r.content)
        announcements = payload.get('announcements') or []
        if any(a.get('secCode') not in {s['code'] for s in stocks} for a in announcements):
            raise ValueError('公告接口未遵守证券过滤')
        ids = [a['announcementId'] for a in announcements]
        if len(set(ids)) != len(ids) or any(i in seen for i in ids):
            raise ValueError('公告分页重复，不能宣称目录完整')
        seen.update(ids)
        found.extend(announcements)
        if len(found) >= expected:
            break
        if not announcements:
            raise ValueError('公告分页中断')
        page += 1
    if len(found) != expected:
        raise ValueError('公告分页记录数与源端总数不一致')
    return found


def extract_amount(kind, text):
    compact = re.sub(r'\s+', '', text).replace('，', ',')
    number = r'([\d,]+(?:\.\d+)?)(亿|万)?元'
    patterns = ([r'共计派发现金红利(?:人民币)?' + number,
                 r'(?:现金红利|现金分红)(?:总额|总计|共计|合计)(?:为|约为|约|人民币|为人民币|:|：)*' + number,
                 r'(?:共计|合计)(?:派发|分配)?现金(?:股利|红利)(?:人民币)?' + number]
                if kind == 'dividend' else
                [r'(?:本次注销[^。；]{0,40})?对应(?:的)?回购(?:金额|资金|成本)(?:为|合计|人民币|为人民币|:|：)*' + number])
    matches = []
    for pattern in patterns:
        for m in re.finditer(pattern, compact):
            if kind == 'buyback_cancelled' and re.search(r'(?<!不)含[^。；]{0,25}(?:交易费用|佣金|印花税)', compact[m.end():m.end()+60]):
                continue
            amount = Decimal(m.group(1).replace(',', '')) * {'亿': Decimal('100000000'), '万': Decimal('10000'), None: Decimal('1')}[m.group(2)]
            matches.append((amount, m.group(0)))
    values = {m[0] for m in matches}
    if not values and kind == 'buyback_cancelled':
        # Use the explicitly paid total ONLY when the same announcement
        # states that every bought share was cancelled. Never prorate costs.
        bought = {v.replace(',', '') for v in re.findall(r'累计回购(?:公司)?(?:的)?股份(?:数量|总数)?(?:为|合计|共计)?([\d,]+)股', compact)}
        cancelled = {v.replace(',', '') for v in re.findall(r'(?:本次(?:回购)?注销(?:(?:已)?回购)?(?:的)?股份(?:数量|数量合计)?(?:为|合计)?|办理完毕上述)([\d,]+)股', compact)}
        paid = list(re.finditer(r'(?:支付(?:的)?总金额|已使用资金总额)(?:为|人民币|为人民币|:|：)*' + number, compact))
        if len(bought) == len(cancelled) == len(paid) == 1 and bought == cancelled:
            m = paid[0]
            if re.search(r'(?<!不)含[^。；]{0,25}(?:交易费用|佣金|印花税)', compact[m.end():m.end()+60]):
                return None, '总金额包含交易费用且未单列股东取得的金额，不计入'
            amount = Decimal(m.group(1).replace(',', '')) * {'亿': Decimal('100000000'), '万': Decimal('10000'), None: Decimal('1')}[m.group(2)]
            return amount, f'{m.group(0)}；公告累计回购与本次注销均为{next(iter(bought))}股，全数注销，不按比例推算'
    if len(values) != 1:
        return None, '未找到唯一明确的本次实际总金额；不能推算'
    amount = next(iter(values))
    return amount, next(quote for value, quote in matches if value == amount)


def process_pdf(kind, a, folder, cutoff, dividends):
    title = re.sub('<[^>]+>', '', a['announcementTitle'])
    url = 'https://static.cninfo.com.cn/' + a['adjunctUrl'].lstrip('/')
    base = {'股票代码': a['secCode'], '公告标题': title, '公告ID': a['announcementId'], '来源': url,
        '公告日期': a['adjunctUrl'].split('/')[1], '类别': kind,
        '实际总金额（元；逐笔，非完整累计）': None, '是否纳入已取得小计': False, '金额原文': None}
    try:
        path = folder / (a['announcementId'] + '.pdf')
        if path.exists():
            content = path.read_bytes()
        else:
            r = requests.get(url, timeout=35)
            r.raise_for_status()
            content = r.content
        if not content.startswith(b'%PDF'):
            raise ValueError('公告下载未返回PDF')
        base['原始PDF_SHA256'] = hashlib.sha256(content).hexdigest()
        path.write_bytes(content)
        texts = [p.extract_text() or '' for p in PdfReader(BytesIO(content)).pages]
        text = '\n'.join(texts)
        path.with_suffix('.txt').write_text(text, encoding='utf-8')
        amount, quote = extract_amount(kind, text)
        base['金额原文'] = quote
        base['实际总金额（元；逐笔，非完整累计）'] = amount
        if kind == 'dividend':
            notice = date.fromisoformat(base['公告日期'])
            matches = [d for d in dividends if d.get('SECURITY_CODE') == a['secCode'] and d.get('ASSIGN_PROGRESS') == '实施分配' and d.get('NOTICE_DATE') and d.get('EX_DIVIDEND_DATE')
                and abs((date.fromisoformat(d['NOTICE_DATE'][:10]) - notice).days) <= 3
                and date.fromisoformat(d['EX_DIVIDEND_DATE'][:10]) <= cutoff]
            if len(matches) == 1 and amount is not None:
                base['是否纳入已取得小计'] = True
                base['经济事件ID'] = a['secCode'] + ':' + str(matches[0].get('REPORT_DATE')) + ':' + matches[0]['EX_DIVIDEND_DATE']
                base['实施日期'] = matches[0]['EX_DIVIDEND_DATE'][:10]
                base['状态'] = '实施公告明确总额，且与已实施除息记录对应；不证明其他历史完整'
            else:
                base['状态'] = '待核实：金额或实施记录无法唯一对应，不汇总'
        else:
            compact = re.sub(r'\s+', '', text)
            confirmed = any(s in compact for s in ['注销事宜已办理完成', '注销手续已办理完毕', '完成注销', '注销已完成', '注销手续已办理完成']) or bool(re.search(r'(?:回购股份|股份回购|公众股份).*注销完成', title) and not any(s in title for s in ['拟', '计划', '更正']))
            if confirmed and amount is not None and date.fromisoformat(base['公告日期']) <= cutoff:
                base['是否纳入已取得小计'] = True
                base['经济事件ID'] = a['secCode'] + ':' + a['announcementId']
                base['状态'] = '完成公告明确注销及本次对应金额；不证明其他回购历史完整'
            else:
                base['状态'] = '待核实：没有同时找到明确注销完成与对应成本，不汇总'
    except Exception as exc:
        base['状态'] = '下载/解析失败: ' + str(exc)
    return base


def deduplicate(rows):
    """Do not sum conflicting announcements for the same economic event."""
    groups = {}
    for row in rows:
        if row.get('是否纳入已取得小计'):
            groups.setdefault(row['经济事件ID'], []).append(row)
    for group in groups.values():
        if len(group) > 1:
            for row in group:
                row['是否纳入已取得小计'] = False
                row['状态'] = '同一经济事件存在多个公告，待核实更正关系，不重复汇总'
    return rows


def cached_catalogue(stocks, folder):
    records, errors = [], []
    for stock in stocks:
        for key in ('权益分派实施', '回购股份注销完成'):
            found, seen, expected, page = [], set(), None, 1
            try:
                while True:
                    path = folder / f'query_{stock["code"]}_{key}_{page}.json'
                    payload = json.loads(path.read_bytes())
                    count = int(payload['totalAnnouncement'])
                    if expected is None:
                        expected = count
                    elif expected != count:
                        raise ValueError('缓存页总数不一致')
                    rows = payload.get('announcements') or []
                    if any(a['secCode'] != stock['code'] for a in rows):
                        raise ValueError('缓存不是当前发行人的单独查询')
                    ids = [a['announcementId'] for a in rows]
                    if len(set(ids)) != len(ids) or any(i in seen for i in ids):
                        raise ValueError('缓存公告分页重复，保留已有唯一记录但不声明完整')
                    seen.update(ids)
                    found.extend(rows)
                    if len(found) >= expected:
                        if len(found) != expected:
                            raise ValueError('缓存记录数与总数不一致')
                        break
                    if not rows:
                        raise ValueError('缓存页中断')
                    page += 1
            except Exception as exc:
                errors.append({'code': stock['code'], 'keyword': key, 'status': str(exc)})
            records.extend((key, a) for a in found)
    return records, errors


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output', type=Path, default=Path('output/csi300_screen.xlsx'))
    p.add_argument('--as-of', type=date.fromisoformat, default=date.today())
    p.add_argument('--cached', action='store_true')
    p.add_argument('--reparse-cancellations', action='store_true', help='重新核实缓存中的注销公告')
    p.add_argument('--catalog-cache', action='store_true', help='复用已检索的目录，继续下载缺失PDF')
    p.add_argument('--workers', type=int, default=8)
    args = p.parse_args()
    if not 1 <= args.workers <= 12:
        p.error('并发下载线程需为1到12')
    folder = args.output.parent / 'raw_public' / args.as_of.isoformat() / 'official_announcements'
    folder.mkdir(parents=True, exist_ok=True)
    roster = read_roster(folder.parent / 'constituents.xls', args.as_of)
    if args.cached:
        result = read_json((folder / 'parsed.json').read_text())
        if args.reparse_cancellations:
            raw_dividends = read_json((folder.parent / 'snapshot.json').read_text())['datasets']['dividends']
            for i, row in enumerate(result['records']):
                if row['类别'] == 'buyback_cancelled':
                    a = {'secCode': row['股票代码'], 'announcementTitle': row['公告标题'], 'announcementId': row['公告ID'], 'adjunctUrl': row['来源'].split('static.cninfo.com.cn/')[1]}
                    result['records'][i] = process_pdf('buyback_cancelled', a, folder, args.as_of, raw_dividends)
            deduplicate(result['records'])
            (folder / 'parsed.json').write_text(dump(result), encoding='utf-8')
    else:
        if args.catalog_cache:
            payload = json.loads((folder / 'stock_master.json').read_bytes())
        else:
            r = requests.get(MASTER, timeout=35); r.raise_for_status()
            (folder / 'stock_master.json').write_bytes(r.content)
            payload = r.json()
        master = {s['code']: s for s in payload['stockList']}
        missing = [s['code'] for s in roster if s['code'] not in master]
        stocks = [master[s['code']] for s in roster if s['code'] in master]
        # Query one issuer at a time: the portal can repeat page 1 for
        # multi-issuer requests. No assertion of lifetime completeness follows.
        announcements, errors = [], [{'status': '发行人映射缺失', 'code': c} for c in missing]
        with ThreadPoolExecutor(max_workers=4) as pool:
            tasks = {pool.submit(query_batch, [stock], key, period, folder): key for stock in stocks
                for key, period in ([] if args.catalog_cache else [('权益分派实施', f'1990-01-01~{args.as_of}'), ('回购股份注销完成', f'1990-01-01~{args.as_of}') ])}
            for task in as_completed(tasks):
                key = tasks[task]
                try:
                    batch = task.result()
                    announcements.extend((key, a) for a in batch)
                    if len(batch) > 25:
                        print('catalogue', key, len(batch), flush=True)
                except Exception as exc:
                    errors.append({'keyword': key, 'status': str(exc)})
        if args.catalog_cache:
            announcements, cache_errors = cached_catalogue(stocks, folder)
            errors.extend(cache_errors)
        selected, seen = [], set()
        for keyword, a in announcements:
            title = re.sub('<[^>]+>', '', a['announcementTitle'])
            if not str(a.get('adjunctUrl', '')).lower().endswith('.pdf'):
                continue
            if a['announcementId'] in seen:
                continue
            kind = 'dividend' if keyword == '权益分派实施' else 'buyback_cancelled'
            if kind == 'dividend' and (not re.search(r'权益分派实施公告', title) or any(x in title for x in ['更正', '调整', '说明', '英文'])):
                continue
            if kind == 'buyback_cancelled' and any(x in title for x in ['限制性', '激励', '业绩', '补偿', '英文']):
                continue
            seen.add(a['announcementId'])
            selected.append((kind, a))
        raw = read_json((folder.parent / 'snapshot.json').read_text())['datasets']['dividends']
        result = {'records': [], 'errors': errors, 'scope': '检索1990年以来可用的权益分派实施及回购股份注销完成公告；标题遗漏、档案缺失、解析失败和金额缺失仍不保证上市以来完整'}
        with ThreadPoolExecutor(max_workers=args.workers) as pool:
            tasks = [pool.submit(process_pdf, kind, a, folder, args.as_of, raw) for kind, a in selected]
            for task in as_completed(tasks):
                row = task.result()
                result['records'].append(row)
                if len(result['records']) % 40 == 0:
                    print('PDFs', len(result['records']), '/', len(tasks), flush=True)
        deduplicate(result['records'])
        (folder / 'parsed.json').write_text(dump(result), encoding='utf-8')
    book = load_workbook(args.output)
    name = '官方公告逐笔核实'
    if name in book:
        del book[name]
    write_sheet(book, name, result['records'])
    ws = book['已取得数据摘要']
    headers = {cell.value: i for i, cell in enumerate(ws[1], 1)}
    extra = ['官方实施公告已对应现金分红小计（元；非累计）', '官方完成公告已对应注销成本小计（元；非累计）']
    for header in extra:
        if header not in headers:
            idx = ws.max_column + 1; ws.cell(1, idx, header); headers[header] = idx
    count = 0
    for line in range(2, ws.max_row + 1):
        code = ws.cell(line, headers['股票代码']).value
        for kind, header in zip(['dividend', 'buyback_cancelled'], extra):
            ws.cell(line, headers[header]).value = None
            amounts = [Decimal(str(r['实际总金额（元；逐笔，非完整累计）'])) for r in result['records'] if r['股票代码'] == code and r.get('是否纳入已取得小计') and r['类别'] == kind]
            if amounts:
                cell = ws.cell(line, headers[header], format(sum(amounts, Decimal('0')), 'f'))
                cell.data_type = 's'; count += 1
    book['统计说明'].append([result['scope']])
    for error in result['errors']:
        book['统计说明'].append(['官方公告查询异常: ' + dump(error)])
    book.move_sheet('已取得数据摘要', offset=-book.index(book['已取得数据摘要']))
    target = args.output.with_name(args.output.stem + '.official.tmp.xlsx')
    book.save(target); target.replace(args.output)
    print(dump({'official_pdf_records': len(result['records']), 'included_records': sum(r.get('是否纳入已取得小计', False) for r in result['records']), 'summary_money_cells_added': count, 'errors': result['errors']}))


if __name__ == '__main__':
    main()
