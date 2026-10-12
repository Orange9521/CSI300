"""Retrieve auditable public records; observed subtotals are never lifetime totals."""
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
from io import BytesIO
import json
from pathlib import Path

import pandas as pd
import requests
from openpyxl import load_workbook
from screen import SOURCES, read_roster, calculate, export

API = 'https://datacenter-web.eastmoney.com/api/data/v1/get'
REPORTS = {
    'dividends': ('RPT_SHAREBONUS_DET', 'SECURITY_CODE', 'SECURITY_CODE,REPORT_DATE,NOTICE_DATE'),
    'ipo': ('RPTA_APP_IPOAPPLY', 'SECURITY_CODE', 'SECURITY_CODE,LISTING_DATE'),
    'seo': ('RPT_SEO_DETAIL', 'SECURITY_CODE', 'SECURITY_CODE,ISSUE_DATE,FINANCE_CODE'),
    'rights': ('RPT_IPO_ALLOTMENT', 'SECURITY_CODE', 'SECURITY_CODE,LISTING_DATE,FINANCE_CODE'),
    'buybacks': ('RPTA_WEB_GETHGLIST_NEW', 'DIM_SCODE', 'DIM_SCODE,REPURCODE'),
}


def read_json(text):
    return json.loads(text, parse_float=Decimal)


def dump(obj):
    return json.dumps(obj, ensure_ascii=False, default=str, indent=2)


def fetch_chunk(kind, codes, directory):
    report, field, sort = REPORTS[kind]
    params = {'reportName': report, 'columns': 'ALL', 'pageSize': 500, 'pageNumber': 1,
              'sortColumns': sort, 'sortTypes': ','.join(['1'] * len(sort.split(','))),
              'filter': '(' + field + ' in (' + ','.join('"' + c + '"' for c in codes) + '))'}
    collected, diagnostics = [], []
    expected = None
    pages = 1
    page = 1
    while page <= pages:
        params['pageNumber'] = page
        r = requests.get(API, params=params, timeout=35)
        r.raise_for_status()
        payload = read_json(r.text)
        raw = r.content
        digest = hashlib.sha256(raw).hexdigest()
        cache = directory / f'{kind}_{codes[0]}_{page}.json'
        cache.write_bytes(raw)
        meta = {'kind': kind, 'codes': codes, 'page': page, 'url': r.url,
                'retrieved_at_utc': datetime.now(timezone.utc).isoformat(), 'sha256': digest, 'cache': str(cache)}
        result = payload.get('result')
        if not payload.get('success') or not result:
            if payload.get('code') == 9201 and payload.get('message') == '返回数据为空':
                meta.update(status='接口未返回记录；不能据此认定历史为零', count=0)
                diagnostics.append(meta)
                return collected, diagnostics
            raise ValueError(f'{report}: {payload.get("message")}')
        if expected is None:
            expected = int(result['count'])
            pages = int(result['pages'])
        elif int(result['count']) != expected:
            raise ValueError('分页期间记录数量变化，需要重新取得一致快照')
        rows = result['data']
        if any(str(row.get(field)).zfill(6) not in codes for row in rows):
            raise ValueError('接口未遵守成分股过滤条件')
        for row in rows:
            row = dict(row)
            row['_source'] = r.url
            row['_raw_sha256'] = digest
            row['_retrieved_at_utc'] = meta['retrieved_at_utc']
            collected.append(row)
        meta.update(status='成功', count=len(rows), expected_count=expected)
        diagnostics.append(meta)
        page += 1
    if len(collected) != expected:
        raise ValueError('分页行数与源端总数不一致')
    return collected, diagnostics


def decimal(value):
    if value is None:
        return None
    result = Decimal(str(value))
    return result if result.is_finite() and result >= 0 else None


def eligible(row, field, cutoff):
    value = row.get(field)
    return bool(value and date.fromisoformat(str(value)[:10]) <= cutoff)


def financing_amount(kind, row):
    """Source money units vary. Validate IPO units against issue shares and price."""
    raw = decimal(row.get('TOTAL_RAISE_FUNDS'))
    if raw is None:
        return None, '源端募集资金金额缺失，不从计划金额推算'
    if kind != 'ipo':
        return raw, '源端TOTAL_RAISE_FUNDS，人民币元；尚未经公告复核'
    shares, price = decimal(row.get('ISSUE_NUM')), decimal(row.get('ISSUE_PRICE'))
    transferred = decimal(row.get('ORIG_TRANSFER_NUM'))
    if transferred is not None and transferred > 0:
        return None, '存在老股转让，尚未核实发行人实际募集额，排除'
    amount = raw * Decimal('100000000')  # source IPO report uses 亿元, shares 万股
    if shares is None or price is None or abs(amount - shares * Decimal('10000') * price) > max(Decimal('1'), amount * Decimal('0.000001')):
        return None, 'IPO金额单位/发行股数×价格校验未通过，排除'
    return amount, '源端TOTAL_RAISE_FUNDS亿元×1亿；通过发行万股×1万×发行价校验；尚未经公告复核'


def summarize(roster, datasets, cutoff):
    summaries, finance_rows, dividend_rows, buyback_rows = [], [], [], []
    for s in roster:
        c = s['code']
        f = {k: [r for r in datasets[k] if str(r.get('DIM_SCODE' if k == 'buybacks' else 'SECURITY_CODE')).zfill(6) == c] for k in REPORTS}
        valid_div = [r for r in f['dividends'] if r.get('ASSIGN_PROGRESS') == '实施分配' and eligible(r, 'EX_DIVIDEND_DATE', cutoff)]
        amounts = {}
        sources = set()
        flags = []
        for kind, date_field in [('ipo', 'LISTING_DATE'), ('seo', 'ISSUE_LISTING_DATE'), ('rights', 'LISTING_DATE')]:
            found = []
            seen = set()
            for r in f[kind]:
                if not eligible(r, date_field, cutoff):
                    continue
                key = str(r.get('FINANCE_CODE') or (r.get('SECUCODE'), r.get(date_field), r.get('TOTAL_RAISE_FUNDS')))
                if key in seen:
                    raise ValueError(f'重复融资事件: {c}, {kind}, {key}')
                seen.add(key)
                amount, reason = financing_amount(kind, r)
                if amount is not None:
                    found.append(amount)
                else:
                    flags.append(kind + ': ' + reason)
                sources.add(r['_source'])
                finance_rows.append({'股票代码': c, '股票名称': s['name'], '类别': kind, '日期': str(r.get(date_field))[:10],
                    '事件ID': key, '已取得募集金额（元；源端记录）': amount, '源端金额': r.get('TOTAL_RAISE_FUNDS'),
                    '金额口径与限制': reason, '来源': r['_source'], '原始响应SHA256': r['_raw_sha256']})
            amounts[kind] = sum(found, Decimal('0')) if found else None
        for r in f['dividends']:
            included = r in valid_div
            dividend_rows.append({'股票代码': c, '股票名称': s['name'], '报告期': r.get('REPORT_DATE'),
                '股权登记日': r.get('EQUITY_RECORD_DATE'), '除权除息日': r.get('EX_DIVIDEND_DATE'),
                '进度': r.get('ASSIGN_PROGRESS'), '每10股税前派息（元）': r.get('PRETAX_BONUS_RMB'),
                '源端总股本（股；非确认分红基数）': r.get('TOTAL_SHARES'), '方案': r.get('IMPL_PLAN_PROFILE'),
                '实际现金分红总额（元）': None, '状态': '已实施记录，现金总额待公告核实' if included else '未实施/截止日后/日期缺失，不计',
                '来源': r['_source'], '原始响应SHA256': r['_raw_sha256']})
            sources.add(r['_source'])
        for r in f['buybacks']:
            buyback_rows.append({'股票代码': c, '股票名称': s['name'], '回购事件ID': r.get('REPURCODE'),
                '进度代码': r.get('REPURPROGRESS'), '源端已回购金额（元；不计股东回报）': r.get('REPURAMOUNT'),
                '源端已回购股数': r.get('REPURNUM'), '完成日': r.get('FINISHDATE'), '公告日期': r.get('UPDATEDATE'),
                '回购目的': r.get('REPUROBJECTIVE'), '备注': r.get('REMARK'),
                '已核实注销金额（元）': None, '状态': '未取得注销完成及对应成本凭证，不计入股东回报',
                '来源': r['_source'], '原始响应SHA256': r['_raw_sha256']})
            sources.add(r['_source'])
        observed = [v for v in amounts.values() if v is not None]
        summaries.append({'股票代码': c, '股票名称': s['name'], '成分股日期': s['date'], '统计截止日': cutoff.isoformat(),
            '已取得IPO募集金额（元；源端记录）': amounts['ipo'], '已取得增发募集金额小计（元）': amounts['seo'],
            '已取得配股募集金额小计（元）': amounts['rights'],
            '已取得三类融资小计（元；非完整股权融资）': sum(observed, Decimal('0')) if observed else None,
            '已实施分红记录数': len(valid_div), '最早已取得分红除息日': min((r['EX_DIVIDEND_DATE'][:10] for r in valid_div), default=None),
            '最近已取得分红除息日': max((r['EX_DIVIDEND_DATE'][:10] for r in valid_div), default=None),
            '已取得回购记录数（非已注销数）': len(f['buybacks']),
            '累计现金分红（元）': None, '累计回购注销（元）': None, '累计股权融资（元）': None,
            '净股东回报（元）': None, '股东回报/融资比率': None,
            '状态': '源端记录已取得；完整历史、现金分红总额、注销及其他股权融资尚未核实',
            '异常': '; '.join(flags), '来源': '\n'.join(sorted(sources))})
    return summaries, finance_rows, dividend_rows, buyback_rows


def write_sheet(book, name, rows):
    ws = book.create_sheet(name)
    if not rows:
        ws.append(['状态']); ws.append(['未取得记录，不等于金额为0'])
        return
    columns = list(rows[0])
    ws.append(columns)
    for row in rows:
        ws.append([format(v, 'f') if isinstance(v, Decimal) else v for v in (row.get(k) for k in columns)])
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 25


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--as-of', type=date.fromisoformat, default=date.today())
    parser.add_argument('--output', type=Path, default=Path('output/csi300_screen.xlsx'))
    parser.add_argument('--cached', action='store_true', help='仅重新生成最近缓存快照，必须与截止日一致')
    args = parser.parse_args()
    if args.as_of > date.today():
        parser.error('截止日不能晚于今天')
    directory = args.output.parent / 'raw_public' / args.as_of.isoformat()
    directory.mkdir(parents=True, exist_ok=True)
    if args.cached:
        roster = read_roster(directory / 'constituents.xls', args.as_of)
        snapshot = read_json((directory / 'snapshot.json').read_text())
        datasets, diagnostics = snapshot['datasets'], snapshot['diagnostics']
    else:
        r = requests.get(SOURCES['中证指数成分股'], timeout=35)
        r.raise_for_status()
        roster = read_roster(BytesIO(r.content), args.as_of)
        (directory / 'constituents.xls').write_bytes(r.content)
        datasets = {k: [] for k in REPORTS}
        diagnostics = [{'kind': 'constituents', 'status': '成功', 'url': r.url, 'count': len(roster), 'data_date': roster[0]['date']}]
        codes = [s['code'] for s in roster]
        with ThreadPoolExecutor(max_workers=4) as pool:
            jobs = {pool.submit(fetch_chunk, k, codes[i:i+50], directory): (k, codes[i:i+50]) for k in REPORTS for i in range(0, len(codes), 50)}
            for job in as_completed(jobs):
                kind, group = jobs[job]
                try:
                    rows, meta = job.result()
                    datasets[kind].extend(rows)
                    diagnostics.extend(meta)
                    print(kind, group[0], len(rows), flush=True)
                except Exception as exc:
                    diagnostics.append({'kind': kind, 'codes': group, 'status': '失败: ' + str(exc)})
                    print(kind, group[0], 'FAILED', str(exc), flush=True)
        (directory / 'snapshot.json').write_text(dump({'datasets': datasets, 'diagnostics': diagnostics}), encoding='utf-8')
    summary, finance, dividends, buybacks = summarize(roster, datasets, args.as_of)
    notes = ['本文件包含真实取得的源端逐笔记录与小计，不是已完成的上市以来累计筛选。',
        '融资小计仅涵盖接口返回的IPO、增发、配股；优先股、跨市场发行、转股、非现金发行及遗漏未核实。',
        '源端分红总股本未证实为实际分红基数，不相乘冒充总金额。',
        '源端已回购金额不是已注销金额，不计入股东回报。',
        '接口分页完整不等于公司历史完整；严格筛选仍要求独立完整性核实。',
        f'统计截止日{args.as_of}；成分股日期{roster[0]["date"]}；检索时间见逐笔原始缓存。',
        '金额采用Decimal，Excel保留十进制文本，缺失保持空白。',
        '原始响应及SHA256留存在output/raw_public，可重复生成；请求失败见数据源检查。']
    empty_ledger = {'events': [], 'coverage': []}
    all_rows, chosen = calculate(roster, empty_ledger, args.as_of)
    # Atomic publication after validating the replacement workbook.
    target = args.output.with_name(args.output.stem + '.tmp.xlsx')
    export(target, all_rows, chosen, [{'类别': d.get('kind'), '状态': d.get('status'), '记录数': d.get('count'), '来源': d.get('url'), '校验SHA256': d.get('sha256')} for d in diagnostics], notes, empty_ledger)
    book = load_workbook(target)
    for name, rows in [('已取得数据摘要', summary), ('融资逐笔记录', finance), ('分红逐笔记录', dividends), ('回购待核实记录', buybacks)]:
        write_sheet(book, name, rows)
    book.move_sheet('已取得数据摘要', offset=-book.index(book['已取得数据摘要']))
    book.save(target)
    checked = load_workbook(target, read_only=True)
    assert checked['已取得数据摘要'].max_row == 301
    checked.close()
    target.replace(args.output)
    print(dump({'output': str(args.output), 'roster_date': roster[0]['date'], 'rows': {k: len(v) for k, v in datasets.items()},
        'companies_with_observed_financing': sum(r['已取得三类融资小计（元；非完整股权融资）'] is not None for r in summary),
        'source_failures': sum(str(d.get('status', '')).startswith('失败') for d in diagnostics), 'strict_screened_companies': len(chosen)}))


if __name__ == '__main__':
    main()
